import json
from io import BytesIO
from pathlib import Path

from docx import Document

TESTS = Path("data/cache/tests.json")
REQS = Path("data/cache/requirements.json")


def _as_list(v):
    if v is None:
        return []
    return v if isinstance(v, list) else [str(v)]


def _as_text(v):
    if v is None:
        return ""
    return ", ".join(map(str, v)) if isinstance(v, list) else str(v)


def build_test_package_bytes() -> bytes:
    tests = json.loads(TESTS.read_text())
    reqs = {r["clause_id"]: r for r in json.loads(REQS.read_text())}
    doc = Document()
    doc.add_heading("Test package", level=0)
    doc.add_paragraph(f"{len(tests)} tests. Drafts are not approved until signed off.")
    for t in tests:
        r = reqs.get(t["clause_id"], {})
        doc.add_heading(f'{t["id"]} ({t["level"]}, Tier {t["tier"]})', level=2)
        doc.add_paragraph(f'Clause {t["clause_id"]}: {r.get("text", "")}')
        doc.add_paragraph(f'Status: {t["status"]}')
        doc.add_paragraph("Steps:")
        for s in _as_list(t.get("steps")):
            doc.add_paragraph(s, style="List Number")
        doc.add_paragraph("Instruments: " + _as_text(t.get("instruments")))
        doc.add_paragraph("Acceptance: " + _as_text(t.get("acceptance")))
        doc.add_paragraph("Reason: " + _as_text(t.get("reason")))
        doc.add_paragraph("Sign-off: ____________________   Date: __________")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
